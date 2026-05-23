from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader, UnstructuredHTMLLoader, CSVLoader, TextLoader
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from typing import List
from langchain_core.documents import Document
import os
import re
import csv
from dotenv import load_dotenv
load_dotenv()


GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
EMBEDDING_PROVIDER = os.getenv("EMBEDDING_PROVIDER", "auto").lower()
GOOGLE_EMBEDDING_MODEL = os.getenv("GOOGLE_EMBEDDING_MODEL", "models/gemini-embedding-001")
HF_EMBEDDING_MODEL = os.getenv("HF_EMBEDDING_MODEL", "sentence-transformers/all-MiniLM-L6-v2")
BASE_DIR = os.path.dirname(os.path.abspath(__file__))


def _build_embedding_function():
    if EMBEDDING_PROVIDER not in {"auto", "google", "huggingface"}:
        raise ValueError("EMBEDDING_PROVIDER must be one of: auto, google, huggingface")

    if EMBEDDING_PROVIDER in {"auto", "google"}:
        if not GOOGLE_API_KEY and EMBEDDING_PROVIDER == "google":
            raise ValueError("GOOGLE_API_KEY environment variable not set")

        if GOOGLE_API_KEY:
            try:
                google_embeddings = GoogleGenerativeAIEmbeddings(
                    model=GOOGLE_EMBEDDING_MODEL,
                    google_api_key=GOOGLE_API_KEY,
                )
                # Probe once so invalid model/key issues fail fast and can fallback in auto mode.
                google_embeddings.embed_query("embedding health check")
                print(f"Using Google embeddings model: {GOOGLE_EMBEDDING_MODEL}")
                return google_embeddings, "google"
            except Exception as e:
                if EMBEDDING_PROVIDER == "google":
                    raise
                print(f"Falling back to HuggingFace embeddings because Google embeddings failed: {e}")

    hf_embeddings = HuggingFaceEmbeddings(model_name=HF_EMBEDDING_MODEL)
    print(f"Using HuggingFace embeddings model: {HF_EMBEDDING_MODEL}")
    return hf_embeddings, "huggingface"

text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200, length_function=len)
embedding_function, resolved_provider = _build_embedding_function()
persist_directory = os.getenv("CHROMA_PERSIST_DIRECTORY", os.path.join(BASE_DIR, f"chroma_db_{resolved_provider}"))
vectorstore = Chroma(persist_directory=persist_directory, embedding_function=embedding_function)


def _is_google_quota_error(error: Exception) -> bool:
    message = str(error).lower()
    return (
        "quota exceeded" in message
        or "rate limit" in message
        or ("429" in message and "embed" in message)
    )


def _switch_to_huggingface_vectorstore() -> None:
    global embedding_function, resolved_provider, persist_directory, vectorstore

    if resolved_provider == "huggingface":
        return

    embedding_function = HuggingFaceEmbeddings(model_name=HF_EMBEDDING_MODEL)
    resolved_provider = "huggingface"
    persist_directory = os.getenv("CHROMA_PERSIST_DIRECTORY_HUGGINGFACE", os.path.join(BASE_DIR, "chroma_db_huggingface"))
    vectorstore = Chroma(persist_directory=persist_directory, embedding_function=embedding_function)
    print(f"Switched to HuggingFace embeddings model: {HF_EMBEDDING_MODEL}")

def load_and_split_document(file_path: str) -> List[Document]:
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        loader = PyPDFLoader(file_path)
        documents = loader.load()
    elif ext == '.docx':
        loader = Docx2txtLoader(file_path)
        documents = loader.load()
    elif ext == '.html':
        try:
            loader = UnstructuredHTMLLoader(file_path)
            documents = loader.load()
        except Exception:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                html_content = f.read()
            text_content = re.sub(r"<[^>]+>", " ", html_content)
            text_content = re.sub(r"\s+", " ", text_content).strip()
            documents = [Document(page_content=text_content, metadata={"source": file_path})]
    elif ext == '.txt':
        loader = TextLoader(file_path, encoding="utf-8")
        documents = loader.load()
    elif ext == '.csv':
        try:
            loader = CSVLoader(file_path, encoding="utf-8")
            documents = loader.load()
        except Exception:
            # Fallback: read CSV as plain text with row context
            documents = []
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                headers = next(reader, None)
                header_str = ", ".join(headers) if headers else ""
                for i, row in enumerate(reader):
                    row_text = ", ".join(f"{h}: {v}" for h, v in zip(headers, row)) if headers else ", ".join(row)
                    documents.append(Document(
                        page_content=row_text,
                        metadata={"source": file_path, "row": i + 1}
                    ))
    elif ext == '.pptx':
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            documents = []
            for slide_num, slide in enumerate(prs.slides, start=1):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = paragraph.text.strip()
                            if text:
                                texts.append(text)
                if texts:
                    documents.append(Document(
                        page_content="\n".join(texts),
                        metadata={"source": file_path, "slide": slide_num}
                    ))
        except ImportError:
            raise ValueError("python-pptx is required for .pptx files. Install it with: pip install python-pptx")
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    # Some PDFs (especially scanned/image-only files) may load with empty text.
    non_empty_docs = [doc for doc in documents if doc.page_content and doc.page_content.strip()]
    if not non_empty_docs:
        raise ValueError("No extractable text found in document. If this is a scanned PDF, run OCR first.")

    splits = text_splitter.split_documents(non_empty_docs)
    non_empty_splits = [split for split in splits if split.page_content and split.page_content.strip()]
    if not non_empty_splits:
        raise ValueError("No usable text chunks found after splitting. Try a text-based PDF or OCR.")

    return non_empty_splits

def index_document_to_chroma(file_path: str, file_id: int) -> tuple[bool, str | None]:
    try:
        splits = load_and_split_document(file_path)
        
        # Add metadata to each split
        for split in splits:
            split.metadata['file_id'] = file_id
        
        try:
            vectorstore.add_documents(splits)
            # vectorstore.persist()
            return True, None
        except Exception as add_error:
            if resolved_provider == "google" and _is_google_quota_error(add_error):
                print("Google embedding quota/rate limit detected. Retrying with HuggingFace embeddings.")
                _switch_to_huggingface_vectorstore()
                vectorstore.add_documents(splits)
                return True, None
            raise
    except Exception as e:
        error_message = f"{type(e).__name__}: {e}"
        print(f"Error indexing document: {error_message}")
        return False, error_message

def delete_doc_from_chroma(file_id: int):
    try:
        docs = vectorstore.get(where={"file_id": file_id})
        print(f"Found {len(docs['ids'])} document chunks for file_id {file_id}")
        
        vectorstore._collection.delete(where={"file_id": file_id})
        print(f"Deleted all documents with file_id {file_id}")
        
        return True
    except Exception as e:
        print(f"Error deleting document with file_id {file_id} from Chroma: {str(e)}")
        return False
